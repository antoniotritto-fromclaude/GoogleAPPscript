#!/usr/bin/env python3
"""Script per estrarre dati dal file Excel PIPELINE"""

import openpyxl
import json
import sys

def extract_data(excel_path):
    wb = openpyxl.load_workbook(excel_path)

    # Database Clienti
    sheet = wb['📊 DATABASE CLIENTI']
    clienti = []

    for row in range(3, min(sheet.max_row + 1, 100)):  # Limita a 100 per test
        nome = sheet.cell(row=row, column=2).value
        if not nome or nome == 'TOT':
            continue

        cliente = {
            'id': sheet.cell(row=row, column=1).value,
            'nome': str(nome).strip(),
            'tipo_cliente': sheet.cell(row=row, column=3).value or 'Potenziale',
            'fonte': sheet.cell(row=row, column=4).value,
            'data_primo_contatto': str(sheet.cell(row=row, column=5).value)[:10] if sheet.cell(row=row, column=5).value else None,
            'stadio_pipeline': sheet.cell(row=row, column=6).value or 'Prospect',
            'probabilita': sheet.cell(row=row, column=7).value or 10,
            'somma_potenziale': sheet.cell(row=row, column=8).value or 0,
            'data_versamento': str(sheet.cell(row=row, column=9).value)[:10] if sheet.cell(row=row, column=9).value else None,
            'somma_versata': sheet.cell(row=row, column=10).value or 0,
            'stato_contabilita': sheet.cell(row=row, column=11).value,
            'ultimo_contatto_tipo': sheet.cell(row=row, column=12).value,
            'data_ultimo_contatto': str(sheet.cell(row=row, column=13).value)[:10] if sheet.cell(row=row, column=13).value else None,
            'prossima_azione': sheet.cell(row=row, column=14).value,
            'note': sheet.cell(row=row, column=15).value,
            'cluster_cliente': sheet.cell(row=row, column=16).value,
            'tipo_fee': sheet.cell(row=row, column=17).value or 'Fondo',
            'funnel_status': sheet.cell(row=row, column=20).value or 'Non avviato',
            'email': sheet.cell(row=row, column=22).value
        }

        # Converti valori numerici
        try:
            cliente['probabilita'] = int(float(cliente['probabilita'])) if cliente['probabilita'] else 10
        except:
            cliente['probabilita'] = 10

        try:
            cliente['somma_potenziale'] = float(cliente['somma_potenziale']) if cliente['somma_potenziale'] else 0
        except:
            cliente['somma_potenziale'] = 0

        try:
            cliente['somma_versata'] = float(cliente['somma_versata']) if cliente['somma_versata'] else 0
        except:
            cliente['somma_versata'] = 0

        clienti.append(cliente)

    # Template Email
    sheet_email = wb['📧 TEMPLATE EMAIL']
    templates = []

    for row in range(3, 15):
        cluster = sheet_email.cell(row=row, column=1).value
        step = sheet_email.cell(row=row, column=2).value
        if not cluster or not step:
            continue

        template = {
            'cluster': cluster,
            'step': int(float(step)),
            'giorni_attesa': int(float(sheet_email.cell(row=row, column=3).value or 0)),
            'oggetto': sheet_email.cell(row=row, column=4).value,
            'corpo': sheet_email.cell(row=row, column=5).value,
            'link_cta': sheet_email.cell(row=row, column=6).value
        }
        templates.append(template)

    # Log Invii
    sheet_log = wb['📬 LOG INVII']
    log_invii = []

    for row in range(2, sheet_log.max_row + 1):
        data = sheet_log.cell(row=row, column=1).value
        if not data:
            continue

        log = {
            'data_invio': str(data)[:19] if data else None,
            'nome': sheet_log.cell(row=row, column=2).value,
            'cluster': sheet_log.cell(row=row, column=3).value,
            'step': sheet_log.cell(row=row, column=4).value,
            'oggetto': sheet_log.cell(row=row, column=5).value,
            'email': sheet_log.cell(row=row, column=6).value,
            'esito': sheet_log.cell(row=row, column=7).value or 'OK'
        }
        log_invii.append(log)

    return {
        'clienti': clienti,
        'templates': templates,
        'log_invii': log_invii
    }

if __name__ == '__main__':
    excel_path = sys.argv[1] if len(sys.argv) > 1 else '../../../PIPELINE_Antonio_Tritto_2026.xlsx'
    data = extract_data(excel_path)
    print(json.dumps(data, ensure_ascii=False))
